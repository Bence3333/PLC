# -*- coding: utf-8 -*-
"""
JEGYZŐKÖNYV-KÉSZÍTŐ - Word (.docx) és PDF jegyzőkönyv a RoboDK-ban rögzített lépésekből és mérésekből
====================================================================================================

A robodk_scripts/lepes_rogzitese.py és robodk_scripts/meresek_rogzitese.py által készített
"<állomás>_dokumentacio" mappákból fejezetekre bontott jegyzőkönyvet készít: címlap,
tartalomjegyzék, feladatonként leírás, környezet, cellafelépítés, lépések képekkel, célpontok,
programok, mérési eredmények (táblázatok + grafikonok) és automatikus értékelés. A sárgával
kiemelt [KITÖLTENDŐ] helyekre a saját szöveged kerül.

Telepítés (egyszer):
    python -m pip install python-docx matplotlib

Használat (egy vagy több feladat mappája, sorrendben):
    python jegyzokonyv_keszito.py "D:/RoboDK/Feladat1_dokumentacio" "D:/RoboDK/Feladat2_dokumentacio" ^
        --nev "Minta Péter" --neptun ABC123 --targy "Ipari robotok" --pdf

Tipp: tegyél minden adatmappába egy feladat.txt fájlt (UTF-8): az első sora a feladat címe,
a többi a feladat szövege - így ez is bekerül a jegyzőkönyvbe.
"""

import argparse
import csv
import datetime
import glob
import json
import math
import os
import pathlib
import shutil
import subprocess
import sys
import tempfile

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

NBSP = "\u00a0"
HONAPOK = ["január", "február", "március", "április", "május", "június", "július", "augusztus",
           "szeptember", "október", "november", "december"]

# Grafikon-színek (egy sorozat: kék; halvány sáv: ugyanannak a skálának világos foka; jelzés: piros)
SZIN = {
    "sorozat": "#2a78d6", "savhatter": "#cde2fb", "kritikus": "#d03b3b", "szoveg": "#0b0b0b",
    "szoveg2": "#52514e", "halvany": "#898781", "racs": "#e1e0d9", "tengely": "#c3c2b7",
}
TABLA_FEJLEC_HATTER = "E8E7E2"
SZOVEG_SZIN = RGBColor(0x1F, 0x23, 0x2B)
HALVANY_SZIN = RGBColor(0x52, 0x51, 0x4E)


# ---------------------------------------------------------------------------
# Számformázás (magyar: tizedesvessző, ezres tagolás szóközzel)
# ---------------------------------------------------------------------------
def sz(v, tizedes=1, egyseg=""):
    if v is None:
        return "–"
    try:
        f = float(v)
    except (TypeError, ValueError):
        return str(v)
    if math.isnan(f) or math.isinf(f):
        return "–"
    s = "{:,.{t}f}".format(abs(f), t=tizedes).replace(",", NBSP).replace(".", ",")
    if f < 0 and s.strip("0,\u00a0"):
        s = "-" + s
    if not egyseg:
        return s
    return s + (egyseg if egyseg in ("°", "%") else NBSP + egyseg)


def szazalek(v, tizedes=1):
    return "–" if v is None else sz(v, tizedes) + "%"


def png_meret(fajl):
    """PNG kép szélessége és magassága pixelben (a fájl fejlécéből)."""
    import struct
    with open(fajl, "rb") as f:
        fej = f.read(24)
    if fej[:8] != b"\x89PNG\r\n\x1a\n":
        return None
    return struct.unpack(">II", fej[16:24])


def magyar_datum(d):
    return "%d. %s %d." % (d.year, HONAPOK[d.month - 1], d.day)


# ---------------------------------------------------------------------------
# Adatok beolvasása
# ---------------------------------------------------------------------------
def feladat_adatai(mappa):
    """lepesek.json + a legutóbbi meres_*/meresek.json + feladat.txt beolvasása."""
    eredmeny = {"mappa": mappa, "lepesek": [], "meres": None, "meres_mappa": None, "cim": None, "leiras": [],
                "osszes_meres": []}
    lepes_fajl = os.path.join(mappa, "lepesek.json")
    if os.path.isfile(lepes_fajl):
        with open(lepes_fajl, encoding="utf-8") as f:
            eredmeny["lepesek"] = json.load(f)
    meresek = sorted(d for d in glob.glob(os.path.join(mappa, "meres_*")) if os.path.isfile(os.path.join(d, "meresek.json")))
    for d in meresek:
        with open(os.path.join(d, "meresek.json"), encoding="utf-8") as f:
            eredmeny["osszes_meres"].append({"mappa": d, "adat": json.load(f)})
    if meresek:  # a részletes kiértékelés a legutolsó mérésről készül
        eredmeny["meres_mappa"] = meresek[-1]
        eredmeny["meres"] = eredmeny["osszes_meres"][-1]["adat"]
    feladat_fajl = os.path.join(mappa, "feladat.txt")
    if os.path.isfile(feladat_fajl):
        with open(feladat_fajl, encoding="utf-8-sig") as f:
            sorok = [s.rstrip() for s in f.read().splitlines()]
        if sorok:
            eredmeny["cim"] = sorok[0].strip()
            eredmeny["leiras"] = [s.strip() for s in sorok[1:] if s.strip()]
    if not eredmeny["cim"]:
        nev = (eredmeny["meres"] or {}).get("allomas", {}).get("nev") or os.path.basename(mappa.rstrip("/\\"))
        eredmeny["cim"] = nev.replace("_dokumentacio", "").replace("_", " ")
    return eredmeny


def palya_csv(fajl):
    """palya_*.csv (pontosvessző, tizedesvessző) -> {oszlopnév: [értékek]}"""
    with open(fajl, encoding="utf-8-sig", newline="") as f:
        r = csv.reader(f, delimiter=";")
        fejlec = next(r)
        oszlopok = {h: [] for h in fejlec}
        for sor in r:
            for h, v in zip(fejlec, sor):
                oszlopok[h].append(float(v.replace(",", ".")) if v.strip() else float("nan"))
    return oszlopok


# ---------------------------------------------------------------------------
# Grafikonok (matplotlib)
# ---------------------------------------------------------------------------
_plt = None


def grafikon_modul():
    global _plt
    if _plt is None:
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
        from matplotlib import font_manager
        elerheto = {f.name for f in font_manager.fontManager.ttflist}
        betu = next((b for b in ("Calibri", "Carlito", "Segoe UI", "Arial", "Liberation Sans") if b in elerheto), "DejaVu Sans")
        plt.rcParams.update({
            "font.family": betu, "font.size": 8.5, "axes.titlesize": 9, "axes.titleweight": "bold",
            "axes.titlecolor": SZIN["szoveg"], "axes.labelcolor": SZIN["szoveg2"], "axes.edgecolor": SZIN["tengely"],
            "axes.linewidth": 0.8, "axes.spines.top": False, "axes.spines.right": False, "axes.axisbelow": True,
            "axes.grid": True, "grid.color": SZIN["racs"], "grid.linewidth": 0.6, "grid.linestyle": "-",
            "xtick.color": SZIN["tengely"], "ytick.color": SZIN["tengely"],
            "xtick.labelcolor": SZIN["szoveg2"], "ytick.labelcolor": SZIN["szoveg2"],
            "lines.linewidth": 1.6, "lines.solid_capstyle": "round", "lines.solid_joinstyle": "round",
            "figure.facecolor": "white", "axes.facecolor": "white", "savefig.dpi": 200, "legend.frameon": False,
        })
        _plt = plt
    return _plt


def _vesszos_tengely(ax, x=True, y=True):
    from matplotlib.ticker import FuncFormatter
    form = FuncFormatter(lambda v, _: ("%g" % round(v, 6)).replace(".", ",").replace("-", "\u2212"))
    if x:
        ax.xaxis.set_major_formatter(form)
    if y:
        ax.yaxis.set_major_formatter(form)


def _racs_elrendezes(n):
    oszlop = 2 if n in (2, 4) else min(3, n)
    return oszlop, int(math.ceil(n / float(oszlop)))


def abra_csuklo_idosor(d, n, elotag, mertekegyseg, fajl):
    """Csuklónkénti kis grafikonok (small multiples): érték az idő függvényében."""
    plt = grafikon_modul()
    oszlop, sor = _racs_elrendezes(n)
    fig, tengelyek = plt.subplots(sor, oszlop, figsize=(6.3, 1.75 * sor + 0.35), sharex=True, squeeze=False)
    t = d["ido_s"]
    for k in range(sor * oszlop):
        ax = tengelyek[k // oszlop][k % oszlop]
        if k >= n:
            ax.set_visible(False)
            continue
        ax.plot(t, d[elotag % (k + 1)], color=SZIN["sorozat"])
        ax.set_title("J%d  [%s]" % (k + 1, mertekegyseg), loc="left")
        ax.set_xlim(t[0], t[-1])
        ax.margins(y=0.12)
        _vesszos_tengely(ax)
        if k // oszlop == sor - 1 or k + oszlop >= n:
            ax.set_xlabel("idő [s]")
            ax.tick_params(labelbottom=True)
    fig.tight_layout(h_pad=1.0, w_pad=1.2)
    fig.savefig(fajl, bbox_inches="tight", pad_inches=0.04)
    plt.close(fig)


def abra_tcp_sebesseg(d, fajl):
    plt = grafikon_modul()
    t, v = d["ido_s"], d["TCP_sebesseg_mm_s"]
    fig, ax = plt.subplots(figsize=(6.3, 2.4))
    ax.fill_between(t, v, color=SZIN["sorozat"], alpha=0.10, lw=0)
    ax.plot(t, v, color=SZIN["sorozat"])
    i = max(range(len(v)), key=lambda k: v[k])
    ax.plot([t[i]], [v[i]], "o", ms=6, color=SZIN["sorozat"], mec="white", mew=1.5, zorder=3)
    jobbra = t[i] < t[0] + 0.75 * (t[-1] - t[0])
    ax.annotate("max. " + sz(v[i], 1, "mm/s"), (t[i], v[i]), xytext=(7 if jobbra else -7, 2), textcoords="offset points",
                ha="left" if jobbra else "right", va="bottom", color=SZIN["szoveg2"], fontsize=8)
    ax.set_xlim(t[0], t[-1])
    ax.set_ylim(0, max(v) * 1.15 if max(v) > 0 else 1)
    ax.set_xlabel("idő [s]")
    ax.set_ylabel("TCP sebesség [mm/s]")
    _vesszos_tengely(ax)
    fig.tight_layout()
    fig.savefig(fajl, bbox_inches="tight", pad_inches=0.04)
    plt.close(fig)


def abra_tcp_palya(d, fajl):
    plt = grafikon_modul()
    x, y, z = d["TCP_X_mm"], d["TCP_Y_mm"], d["TCP_Z_mm"]
    fig, (a1, a2) = plt.subplots(1, 2, figsize=(6.3, 3.1))
    zart = math.dist((x[0], y[0], z[0]), (x[-1], y[-1], z[-1])) < 1.0
    for ax, h, cim, cimke in ((a1, y, "Felülnézet (X–Y)", "Y [mm]"), (a2, z, "Oldalnézet (X–Z)", "Z [mm]")):
        ax.plot(x, h, color=SZIN["sorozat"])
        ax.plot([x[0]], [h[0]], "o", ms=7, color=SZIN["sorozat"], mec="white", mew=1.5, zorder=3)
        ax.annotate("kezdet = vég" if zart else "kezdet", (x[0], h[0]), xytext=(6, 6), textcoords="offset points",
                    color=SZIN["szoveg2"], fontsize=8)
        if not zart:
            ax.plot([x[-1]], [h[-1]], "s", ms=6, color=SZIN["sorozat"], mec="white", mew=1.5, zorder=3)
            ax.annotate("vég", (x[-1], h[-1]), xytext=(6, -10), textcoords="offset points", color=SZIN["szoveg2"], fontsize=8)
        ax.set_aspect("equal", adjustable="datalim")
        ax.margins(0.12)
        ax.set_title(cim, loc="left")
        ax.set_xlabel("X [mm]")
        ax.set_ylabel(cimke)
        _vesszos_tengely(ax)
    fig.tight_layout(w_pad=2.0)
    fig.savefig(fajl, bbox_inches="tight", pad_inches=0.04)
    plt.close(fig)


def abra_kihasznaltsag(csuklok, fajl):
    """Csuklónként: teljes tartomány (halvány sáv) és a program által bejárt rész (kék)."""
    plt = grafikon_modul()
    cs = [c for c in csuklok if "also_hatar" in c]
    if not cs:
        return False
    n = len(cs)
    fig, ax = plt.subplots(figsize=(6.3, 0.42 * n + 0.95))
    also = min(c["also_hatar"] for c in cs)
    felso = max(c["felso_hatar"] for c in cs)
    hely = 0.03 * (felso - also)
    for i, c in enumerate(cs):
        y = n - 1 - i
        ax.barh(y, c["felso_hatar"] - c["also_hatar"], left=c["also_hatar"], height=0.42, color=SZIN["savhatter"], lw=0)
        ax.barh(y, max(c["mozgastartomany"], 0.004 * (felso - also)), left=c["min"], height=0.42, color=SZIN["sorozat"], lw=0)
        felirat = szazalek(c.get("kihasznaltsag_szazalek"), 0)
        ax.text(c["felso_hatar"] + hely, y, felirat, va="center", ha="left", color=SZIN["szoveg2"], fontsize=8)
        tartomany = c["felso_hatar"] - c["also_hatar"]
        if tartomany > 0 and min(c["tartalek_also"], c["tartalek_felso"]) < 0.05 * tartomany:
            for dx, szoveg, szin, meret in ((30, "\u26a0", SZIN["kritikus"], 9), (42, "határ közelében", SZIN["szoveg"], 8)):
                ax.annotate(szoveg, (c["felso_hatar"] + hely, y), xytext=(dx, 0), textcoords="offset points",
                            va="center", ha="left", color=szin, fontsize=meret)
    ax.set_yticks(range(n))
    ax.set_yticklabels([c["nev"] for c in reversed(cs)])
    ax.set_xlim(also - hely, felso + (felso - also) * 0.32)
    ax.set_ylim(-0.6, n - 0.4)
    ax.grid(axis="y", visible=False)
    ax.set_xlabel("csuklószög [°]   (világos: teljes tartomány, kék: a program által bejárt rész)")
    _vesszos_tengely(ax, y=False)
    fig.tight_layout()
    fig.savefig(fajl, bbox_inches="tight", pad_inches=0.04)
    plt.close(fig)
    return True


def abra_osszehasonlitas(programok, fajl):
    """Programonként egy panel: a ciklusidő mérésenként (vízszintes oszlopok, érték az oszlop végén)."""
    plt = grafikon_modul()
    programok = programok[:6]
    magassagok = [0.34 * len(p[1]) + 0.55 for p in programok]
    fig, tengelyek = plt.subplots(len(programok), 1, figsize=(6.3, sum(magassagok) + 0.4), squeeze=False,
                                  gridspec_kw={"height_ratios": magassagok})
    for ax, (prog, meresek) in zip([t[0] for t in tengelyek], programok):
        cimkek = [m[0] for m in meresek]
        ertekek = [m[1] for m in meresek]
        y = list(range(len(meresek)))[::-1]
        ax.barh(y, ertekek, height=0.5, color=SZIN["sorozat"], lw=0)
        felso = max(ertekek) if ertekek else 1.0
        for yy, v in zip(y, ertekek):
            ax.text(v + felso * 0.01, yy, sz(v, 2, "s"), va="center", ha="left", color=SZIN["szoveg2"], fontsize=8)
        ax.set_yticks(y)
        ax.set_yticklabels(cimkek)
        ax.set_xlim(0, felso * 1.18)
        ax.set_ylim(-0.6, len(meresek) - 0.4)
        ax.grid(axis="y", visible=False)
        ax.set_title(prog, loc="left")
        _vesszos_tengely(ax, y=False)
    tengelyek[-1][0].set_xlabel("ciklusidő [s]")
    fig.tight_layout(h_pad=1.2)
    fig.savefig(fajl, bbox_inches="tight", pad_inches=0.04)
    plt.close(fig)


def abra_zona_tavolsag(d, fajl):
    """A TCP távolsága a tiltott zónától az idő függvényében."""
    plt = grafikon_modul()
    t, z = d["ido_s"], d["zona_tavolsag_mm"]
    fig, ax = plt.subplots(figsize=(6.3, 2.4))
    ax.axhline(0, color=SZIN["kritikus"], lw=1.0)
    ax.text(t[-1], 0, " zóna határa", va="bottom", ha="right", color=SZIN["szoveg2"], fontsize=8)
    ax.plot(t, z, color=SZIN["sorozat"])
    i = min(range(len(z)), key=lambda k: z[k])
    ax.plot([t[i]], [z[i]], "o", ms=6, color=SZIN["sorozat"], mec="white", mew=1.5, zorder=3)
    jobbra = t[i] < t[0] + 0.75 * (t[-1] - t[0])
    ax.annotate("min. " + sz(z[i], 1, "mm"), (t[i], z[i]), xytext=(7 if jobbra else -7, 4), textcoords="offset points",
                ha="left" if jobbra else "right", va="bottom", color=SZIN["szoveg2"], fontsize=8)
    ax.set_xlim(t[0], t[-1])
    ax.set_ylim(min(0.0, min(z)) - 5, max(z) * 1.1 + 5)
    ax.set_xlabel("idő [s]")
    ax.set_ylabel("távolság a zónától [mm]")
    _vesszos_tengely(ax)
    fig.tight_layout()
    fig.savefig(fajl, bbox_inches="tight", pad_inches=0.04)
    plt.close(fig)


def abra_program_idok(nevek, idok, fajl):
    """Programonként a ciklusidő (vízszintes oszlopok, érték az oszlop végén)."""
    plt = grafikon_modul()
    fig, ax = plt.subplots(figsize=(6.3, 0.34 * len(nevek) + 0.75))
    y = list(range(len(nevek)))[::-1]
    ax.barh(y, idok, height=0.5, color=SZIN["sorozat"], lw=0)
    felso = max(idok) if idok else 1.0
    for yy, v in zip(y, idok):
        ax.text(v + felso * 0.01, yy, sz(v, 2, "s"), va="center", ha="left", color=SZIN["szoveg2"], fontsize=8)
    ax.set_yticks(y)
    ax.set_yticklabels(nevek)
    ax.set_xlim(0, felso * 1.16)
    ax.set_ylim(-0.6, len(nevek) - 0.4)
    ax.grid(axis="y", visible=False)
    ax.set_xlabel("ciklusidő [s]")
    _vesszos_tengely(ax, y=False)
    fig.tight_layout()
    fig.savefig(fajl, bbox_inches="tight", pad_inches=0.04)
    plt.close(fig)


def abra_movej_movel(cimkek, j_idok, l_idok, fajl):
    """MoveJ és MoveL változat ciklusideje változatonként (két sorozat, jelmagyarázattal)."""
    plt = grafikon_modul()
    fig, ax = plt.subplots(figsize=(6.3, 0.62 * len(cimkek) + 0.95))
    y = list(range(len(cimkek)))[::-1]
    h = 0.32
    felso = max(j_idok + l_idok) if (j_idok or l_idok) else 1.0
    for eltolas, idok, szin, nev in ((h / 2 + 0.02, j_idok, SZIN["sorozat"], "MoveJ (csuklómozgás)"),
                                      (-h / 2 - 0.02, l_idok, "#eb6834", "MoveL (lineáris mozgás)")):
        ax.barh([v + eltolas for v in y], idok, height=h, color=szin, lw=0, label=nev)
        for yy, v in zip(y, idok):
            ax.text(v + felso * 0.01, yy + eltolas, sz(v, 2, "s"), va="center", ha="left", color=SZIN["szoveg2"], fontsize=8)
    ax.set_yticks(y)
    ax.set_yticklabels(cimkek)
    ax.set_xlim(0, felso * 1.18)
    ax.grid(axis="y", visible=False)
    ax.set_xlabel("ciklusidő [s]")
    ax.legend(loc="lower right", fontsize=8)
    _vesszos_tengely(ax, y=False)
    fig.tight_layout()
    fig.savefig(fajl, bbox_inches="tight", pad_inches=0.04)
    plt.close(fig)


def abra_kis_tobbszoros(sorozatok, x_kulcs, y_kulcs, x_cimke, y_cimke, fajl, felulnezet=False, zona=False):
    """Programonként egy kis panel azonos skálával (small multiples) - a programok összevetéséhez."""
    plt = grafikon_modul()
    n = len(sorozatok)
    oszlop = 1 if n == 1 else (3 if (felulnezet and n > 4) else 2)
    sor = int(math.ceil(n / float(oszlop)))
    magas = (2.35 if oszlop == 2 else 1.75) if felulnezet else (1.6 if n <= 4 else 1.35)
    fig, tengelyek = plt.subplots(sor, oszlop, figsize=(6.3, magas * sor + 0.35), squeeze=False)
    ymin = min(min(d[y_kulcs]) for _, d in sorozatok)
    ymax = max(max(d[y_kulcs]) for _, d in sorozatok)
    xmin = min(min(d[x_kulcs]) for _, d in sorozatok)
    xmax = max(max(d[x_kulcs]) for _, d in sorozatok)
    for k in range(sor * oszlop):
        ax = tengelyek[k // oszlop][k % oszlop]
        if k >= n:
            ax.set_visible(False)
            continue
        cim, d = sorozatok[k]
        x, y = d[x_kulcs], d[y_kulcs]
        if zona:
            ax.axhline(0, color=SZIN["kritikus"], lw=0.9)
        ax.plot(x, y, color=SZIN["sorozat"], lw=1.3)
        if felulnezet:
            ax.plot([x[0]], [y[0]], "o", ms=5, color=SZIN["sorozat"], mec="white", mew=1.2, zorder=3)
            szel = max(xmax - xmin, ymax - ymin) * 0.55
            kx, ky = (xmin + xmax) / 2, (ymin + ymax) / 2
            ax.set_xlim(kx - szel, kx + szel)
            ax.set_ylim(ky - szel, ky + szel)
            ax.set_aspect("equal", adjustable="box")
        else:
            ax.set_xlim(min(x), max(x))
            ax.set_ylim(min(0.0, ymin) - (5 if zona else 0), ymax * 1.08 + (5 if zona else 0))
            i = max(range(len(y)), key=lambda j: y[j]) if not zona else min(range(len(y)), key=lambda j: y[j])
            ax.plot([x[i]], [y[i]], "o", ms=4.5, color=SZIN["sorozat"], mec="white", mew=1.0, zorder=3)
            ax.annotate(("min. " if zona else "max. ") + sz(y[i], 0 if not zona else 1), (x[i], y[i]),
                        xytext=(5, 3), textcoords="offset points", color=SZIN["szoveg2"], fontsize=7.5,
                        ha="left" if x[i] < min(x) + 0.7 * (max(x) - min(x)) else "right")
        ax.set_title(cim, loc="left", fontsize=8.5)
        ax.tick_params(labelsize=7.5)
        _vesszos_tengely(ax)
        if k % oszlop == 0:
            ax.set_ylabel(y_cimke, fontsize=8)
        if k // oszlop == sor - 1 or k + oszlop >= n:
            ax.set_xlabel(x_cimke, fontsize=8)
    fig.tight_layout(h_pad=0.8, w_pad=1.0)
    fig.savefig(fajl, bbox_inches="tight", pad_inches=0.04)
    plt.close(fig)


def program_alairas(p):
    """Az utasítások szerkezete a számértékek nélkül (azonos szerkezetű programok felismeréséhez)."""
    return tuple((u.get("tipus"), u.get("mozgas"), u["nev"] if u.get("mozgas") else "") for u in p.get("utasitasok", []))


def movej_movel_parok(programok):
    """Olyan programpárok, amelyek neve csak a MoveJ/MoveL szóban tér el: [(változat, MoveJ-prog, MoveL-prog)]."""
    nevek = {p["nev"]: p for p in programok}
    parok = []
    for nev, p in nevek.items():
        if "MoveJ" in nev and nev.replace("MoveJ", "MoveL") in nevek:
            valtozat = nev.replace("MoveJ", "").replace("__", "_").strip("_") or nev
            parok.append((valtozat, p, nevek[nev.replace("MoveJ", "MoveL")]))
    return parok


# ---------------------------------------------------------------------------
# Word segédfüggvények
# ---------------------------------------------------------------------------
def mezo(bekezdes, utasitas, alapertek=""):
    """Word mező (pl. PAGE, NUMPAGES) beszúrása."""
    for tipus, szoveg in (("begin", None), ("instr", utasitas), ("separate", None), ("text", alapertek), ("end", None)):
        run = bekezdes.add_run()
        if tipus == "instr":
            elem = OxmlElement("w:instrText")
            elem.set(qn("xml:space"), "preserve")
            elem.text = szoveg
            run._r.append(elem)
        elif tipus == "text":
            run.text = szoveg
        else:
            elem = OxmlElement("w:fldChar")
            elem.set(qn("w:fldCharType"), tipus)
            run._r.append(elem)


def cella_hatter(cella, szin):
    tcPr = cella._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), szin)
    tcPr.append(shd)


def ismetlodo_fejlec(sor):
    trPr = sor._tr.get_or_add_trPr()
    elem = OxmlElement("w:tblHeader")
    elem.set(qn("w:val"), "true")
    trPr.append(elem)


SETTINGS_KESOBBI = {qn("w:" + t) for t in (
    "hdrShapeDefaults", "footnotePr", "endnotePr", "compat", "docVars", "rsids", "mathPr", "attachedSchema",
    "themeFontLang", "clrSchemeMapping", "doNotIncludeSubdocsInStats", "doNotAutoCompressPictures", "forceUpgrade",
    "captions", "readModeInkLockDown", "smartTagType", "schemaLibrary", "shapeDefaults", "doNotEmbedSmartTags",
    "decimalSymbol", "listSeparator")}


def mezok_frissitese_megnyitaskor(doc):
    """Word megnyitáskor felajánlja a mezők (tartalomjegyzék, oldalszámok) frissítését."""
    beallitasok = doc.settings.element
    zoom = beallitasok.find(qn("w:zoom"))
    if zoom is not None and zoom.get(qn("w:percent")) is None:  # a python-docx sablonjából hiányzik
        zoom.set(qn("w:percent"), "100")
    elem = OxmlElement("w:updateFields")
    elem.set(qn("w:val"), "true")
    kovetkezo = next((c for c in beallitasok if c.tag in SETTINGS_KESOBBI), None)
    if kovetkezo is not None:
        kovetkezo.addprevious(elem)
    else:
        beallitasok.append(elem)


class Jegyzokonyv(object):
    def __init__(self, args):
        self.args = args
        self.doc = Document()
        self.abra_sz = 0
        self.tabla_sz = 0
        self.fejezetek = []  # (szint, cím) a tartalomjegyzékhez
        self._dokumentum_beallitasa()

    # --- alapbeállítások -------------------------------------------------
    def _dokumentum_beallitasa(self):
        a = self.args
        for s in self.doc.sections:
            s.page_width, s.page_height = Cm(21.0), Cm(29.7)  # A4
            s.left_margin = s.right_margin = Cm(2.5)
            s.top_margin, s.bottom_margin = Cm(2.5), Cm(2.2)
            s.header_distance = s.footer_distance = Cm(1.2)
        normal = self.doc.styles["Normal"]
        normal.font.name = a.betutipus
        normal.font.size = Pt(a.betumeret)
        normal.paragraph_format.space_after = Pt(6)
        normal.paragraph_format.line_spacing = 1.15
        rpr = self.doc.styles.element.find(qn("w:docDefaults")).find(qn("w:rPrDefault")).find(qn("w:rPr"))
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.insert(0, rfonts)
        for attr in list(rfonts.attrib):
            if attr.endswith("Theme") or attr.endswith("theme"):
                del rfonts.attrib[attr]
        for attr in ("w:ascii", "w:hAnsi", "w:cs"):
            rfonts.set(qn(attr), a.betutipus)
        nyelv = OxmlElement("w:lang")
        nyelv.set(qn("w:val"), "hu-HU")
        rpr.append(nyelv)
        for szint, meret in ((1, 16), (2, 13), (3, 11.5)):
            st = self.doc.styles["Heading %d" % szint]
            st.font.name = a.betutipus
            st.font.size = Pt(meret)
            st.font.bold = True
            st.font.color.rgb = SZOVEG_SZIN
            st.paragraph_format.space_before = Pt(18 if szint == 1 else 12)
            st.paragraph_format.space_after = Pt(6)
            rfonts = st.element.get_or_add_rPr().find(qn("w:rFonts"))
            for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
                if rfonts is not None and rfonts.get(qn(attr)) is not None:
                    del rfonts.attrib[qn(attr)]
        mezok_frissitese_megnyitaskor(self.doc)

    # --- építőelemek -----------------------------------------------------
    def cimsor(self, szoveg, szint):
        self.doc.add_heading(szoveg, level=szint)
        self.fejezetek.append((szint, szoveg))

    def bekezdes(self, szoveg="", felkover=False, dolt=False, meret=None, szin=None, igazitas=None, utana=None):
        p = self.doc.add_paragraph()
        if szoveg:
            r = p.add_run(szoveg)
            r.bold, r.italic = felkover, dolt
            if meret:
                r.font.size = Pt(meret)
            if szin is not None:
                r.font.color.rgb = szin
        if igazitas is not None:
            p.alignment = igazitas
        if utana is not None:
            p.paragraph_format.space_after = Pt(utana)
        return p

    def kitoltendo(self, szoveg):
        p = self.doc.add_paragraph()
        r = p.add_run("[KITÖLTENDŐ: " + szoveg + "]")
        r.font.highlight_color = WD_COLOR_INDEX.YELLOW
        return p

    def felsorolas(self, szoveg):
        return self.doc.add_paragraph(szoveg, style="List Bullet")

    def megjegyzes(self, szoveg):
        return self.bekezdes(szoveg, dolt=True, meret=self.args.betumeret - 1.5, szin=HALVANY_SZIN)

    def uj_oldal(self):
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def abra(self, kep, cim, szelesseg_cm=16.0):
        if not kep or not os.path.isfile(kep):
            self.kitoltendo("ide kerül a kép („%s”) – mentsd el kézzel ezen a néven: %s, és futtasd újra a "
                            "jegyzőkönyv-készítőt (vagy illeszd be Wordben)." % (cim, kep or "-"))
            return
        meret = png_meret(kep)
        if meret and meret[0] > 0:  # a kép (felirattal együtt) férjen el egy oldalon
            max_magassag_cm = 21.0
            if szelesseg_cm * meret[1] / float(meret[0]) > max_magassag_cm:
                szelesseg_cm = max_magassag_cm * meret[0] / float(meret[1])
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_after = Pt(2)
        p.add_run().add_picture(kep, width=Cm(szelesseg_cm))
        self.abra_sz += 1
        felirat = self.bekezdes("%d. ábra – %s" % (self.abra_sz, cim), dolt=True, meret=self.args.betumeret - 1.5,
                                szin=HALVANY_SZIN, igazitas=WD_ALIGN_PARAGRAPH.CENTER, utana=10)
        return felirat

    def tablazat(self, fejlec, sorok, cim, szelessegek=None, jobbra=None, meret=None):
        """Táblázat felirattal (a táblázat fölött). jobbra: jobbra igazított (szám) oszlopok indexei."""
        meret = meret or self.args.betumeret - 2
        self.tabla_sz += 1
        felirat = self.bekezdes("%d. táblázat – %s" % (self.tabla_sz, cim), dolt=True, meret=self.args.betumeret - 1.5,
                                szin=HALVANY_SZIN, utana=3)
        felirat.paragraph_format.keep_with_next = True
        jobbra = set(jobbra or [])
        t = self.doc.add_table(rows=1, cols=len(fejlec))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        for i, h in enumerate(fejlec):
            c = t.rows[0].cells[i]
            c.text = ""
            r = c.paragraphs[0].add_run(h)
            r.bold = True
            r.font.size = Pt(meret)
            cella_hatter(c, TABLA_FEJLEC_HATTER)
            if i in jobbra:
                c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        ismetlodo_fejlec(t.rows[0])
        for sor in sorok:
            cellak = t.add_row().cells
            for i, ertek in enumerate(sor):
                cellak[i].text = ""
                r = cellak[i].paragraphs[0].add_run(str(ertek))
                r.font.size = Pt(meret)
                if i in jobbra:
                    cellak[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for sor in t.rows:
            for c in sor.cells:
                for p in c.paragraphs:
                    p.paragraph_format.space_after = Pt(0)
                    p.paragraph_format.line_spacing = 1.0
        if szelessegek:
            for i, w in enumerate(szelessegek):
                t.columns[i].width = Cm(w)
                for sor in t.rows:
                    sor.cells[i].width = Cm(w)
        self.bekezdes(utana=4)
        return t

    def kulcs_ertek(self, parok, cim):
        self.tablazat(["Megnevezés", "Érték"], parok, cim, szelessegek=[5.5, 10.5])

    # --- címlap, tartalomjegyzék, fejléc/lábléc ---------------------------
    def cimlap(self, robodk_verzio):
        a = self.args
        self.bekezdes(a.intezmeny, meret=13, szin=HALVANY_SZIN, igazitas=WD_ALIGN_PARAGRAPH.CENTER, utana=120)
        self.bekezdes("JEGYZŐKÖNYV", felkover=True, meret=30, igazitas=WD_ALIGN_PARAGRAPH.CENTER, utana=10)
        self.bekezdes(a.cim, meret=17, igazitas=WD_ALIGN_PARAGRAPH.CENTER, utana=6)
        self.bekezdes(a.targy, meret=13, szin=HALVANY_SZIN, igazitas=WD_ALIGN_PARAGRAPH.CENTER, utana=150)
        sorok = [("Készítette", "%s%s" % (a.nev, " (%s)" % a.neptun if a.neptun else ""))]
        if a.oktato:
            sorok.append(("Oktató", a.oktato))
        sorok.append(("Dátum", magyar_datum(datetime.date.today())))
        if robodk_verzio:
            sorok.append(("Szoftver", "RoboDK " + robodk_verzio))
        for k, v in sorok:
            p = self.bekezdes(igazitas=WD_ALIGN_PARAGRAPH.CENTER, utana=2)
            r = p.add_run(k + ": ")
            r.font.color.rgb = HALVANY_SZIN
            p.add_run(v).bold = True
        if a.minta:
            p = self.bekezdes(igazitas=WD_ALIGN_PARAGRAPH.CENTER)
            p.paragraph_format.space_before = Pt(36)
            r = p.add_run("MINTA – fiktív, szimulált adatokkal készült bemutató, nem valódi RoboDK-mérés. Beadandóként nem használható.")
            r.bold = True
            r.font.color.rgb = RGBColor(0xD0, 0x3B, 0x3B)
        self.uj_oldal()
        self.bekezdes("Tartalomjegyzék", felkover=True, meret=16, utana=10)
        self._toc_hely = self.doc.add_paragraph()
        self.uj_oldal()

    def tartalomjegyzek(self):
        """TOC mező, előre kitöltve a fejezetcímekkel (Wordben F9 / megnyitáskori frissítés adja az oldalszámokat)."""
        hely = self._toc_hely
        bekezdesek = []
        for i, (szint, szoveg) in enumerate(self.fejezetek):
            p = self.doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.6 * (szint - 1))
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.space_before = Pt(6 if szint == 1 else 0)
            if i == 0:
                for tipus, szoveg_ in (("begin", None), ("instr", ' TOC \\o "1-3" \\h \\z \\u '), ("separate", None)):
                    run = p.add_run()
                    if tipus == "instr":
                        e = OxmlElement("w:instrText")
                        e.set(qn("xml:space"), "preserve")
                        e.text = szoveg_
                    else:
                        e = OxmlElement("w:fldChar")
                        e.set(qn("w:fldCharType"), tipus)
                    run._r.append(e)
            r = p.add_run(szoveg)
            r.bold = szint == 1
            bekezdesek.append(p)
        if not bekezdesek:
            return
        vege = OxmlElement("w:fldChar")
        vege.set(qn("w:fldCharType"), "end")
        bekezdesek[-1].add_run()._r.append(vege)
        for p in bekezdesek:
            hely._p.addprevious(p._p)
        megj = self.doc.add_paragraph()
        r = megj.add_run("Az oldalszámokhoz Wordben: jobb klikk a tartalomjegyzékre → Mező frissítése (F9).")
        r.italic = True
        r.font.size = Pt(self.args.betumeret - 2)
        r.font.color.rgb = HALVANY_SZIN
        hely._p.addprevious(megj._p)
        hely._p.getparent().remove(hely._p)

    def fejlec_lablec(self):
        s = self.doc.sections[0]
        s.different_first_page_header_footer = True
        fej = s.header.paragraphs[0]
        r = fej.add_run("%s – %s" % (self.args.cim, self.args.nev))
        r.font.size = Pt(8.5)
        r.font.color.rgb = HALVANY_SZIN
        if self.args.minta:
            r2 = fej.add_run("   |   MINTA – fiktív adatok")
            r2.font.size = Pt(8.5)
            r2.font.color.rgb = RGBColor(0xD0, 0x3B, 0x3B)
        lab = s.footer.paragraphs[0]
        lab.alignment = WD_ALIGN_PARAGRAPH.CENTER
        mezo(lab, " PAGE ", "1")
        lab.add_run(" / ")
        mezo(lab, " NUMPAGES ", "1")
        for run in lab.runs:
            run.font.size = Pt(9)
            run.font.color.rgb = HALVANY_SZIN

    # --- fejezetek -----------------------------------------------------------
    def bevezetes(self, feladatok):
        self.cimsor("Bevezetés", 1)
        verziok = sorted({f["meres"]["allomas"].get("robodk_verzio", "") for f in feladatok if f["meres"]})
        szoveg = ("A jegyzőkönyv %d feladat megoldását dokumentálja. A feladatokat a RoboDK%s robotszimulációs és "
                  "offline programozó szoftverben valósítottam meg."
                  % (len(feladatok), (" " + ", ".join(v for v in verziok if v)) if any(verziok) else ""))
        if any(f["lepesek"] for f in feladatok):
            szoveg += " A fontos munkalépéseket képernyőképpel együtt rögzítettem."
        if any(f["meres"] for f in feladatok):
            szoveg += (" A méréseket (ciklusidő, pályahossz, célpontok koordinátái, csuklószögek, sebességek, "
                       "gyorsulások, csuklóhatároktól való távolság, szingularitás- és ütközésvizsgálat) a RoboDK Python "
                       "API-ján keresztül, automatikusan mentettem ki, így a közölt értékek közvetlenül a szimulációból "
                       "származnak.")
        self.bekezdes(szoveg)
        self.bekezdes("Mértékegységek: hossz mm, szög fok (°), idő s. A koordináták a célpont szülő-referenciakeretéhez "
                      "képest értendők; az orientáció Rx→Ry′→Rz″ sorrendű Euler-szögekkel van megadva "
                      "(H = Transl·Rot(x)·Rot(y)·Rot(z)).")
        self.kitoltendo("a gyakorlat célja néhány mondatban (pl. ipari robot offline programozásának és szimulációjának "
                        "elsajátítása, ciklusidő-becslés, elérhetőség vizsgálata).")
        if any(f["meres"] for f in feladatok):
            self.cimsor("Mérési módszer és korlátai", 2)
            for sor in (
                "Ciklusidő és pályahossz: a RoboDK pályaszámítása (program frissítése) adja. A mérés előtt a robot a "
                "program első célpontjába áll, így az eredmény nem függ attól, hol állt korábban a robot.",
                "Pálya: a RoboDK időalapú mintavételezésével rögzített csuklószögek és TCP-pozíciók (az időlépés "
                "programonként a táblázatokban szerepel). A sebességek és gyorsulások ebből származnak.",
                "Célpontok elérése: a célpontot a TCP helyzete alapján azonosítom (tűrés: 0,1 mm). Az eltérés a célpont "
                "és a mintapontok közötti, interpolált pálya legkisebb távolsága. Megálló célpontnál ez a mintavétel "
                "felbontásán belüli érték, lekerekített (nem megálló) célpontnál a sarok levágásának mértéke.",
                "Korlát: a RoboDK ideális kinematikai modellel szimulál. A megálló célpontokat a robot sebességtől "
                "függetlenül pontosan eléri, és a mozgás determinisztikus, ezért az ismételhetőség is ideális. "
                "A valódi robot pozicionálási pontosságát és ismételhetőségét (ISO 9283) a hajtások dinamikája, "
                "a mechanikai rugalmasság és a tűrések határozzák meg, ezeket a szimuláció nem modellezi.",
            ):
                self.felsorolas(sor)

    def feladat(self, fsz, f):
        self.uj_oldal()
        self.cimsor("%d. feladat – %s" % (fsz, f["cim"]), 1)
        m = f["meres"] or {}
        mm = f["meres_mappa"]

        # 1. leírás
        self.cimsor("%d.1 A feladat leírása" % fsz, 2)
        if f["leiras"]:
            for sor in f["leiras"]:
                if sor[:2] in ("- ", "* ", "• "):
                    self.felsorolas(sor[2:].strip())
                else:
                    self.bekezdes(sor)
        else:
            self.kitoltendo("másold ide a feladat szövegét – vagy írd a „feladat.txt” fájlba ide: %s "
                            "(első sor: cím), és futtasd újra a jegyzőkönyv-készítőt." % f["mappa"])

        # 2. környezet
        self.cimsor("%d.2 A szimulációs környezet" % fsz, 2)
        if not m:
            self.kitoltendo("nem találtam mérési adatot (meres_*/meresek.json) ebben a mappában: %s – futtasd a "
                            "meresek_rogzitese.py szkriptet a RoboDK-ban." % f["mappa"])
        else:
            al = m.get("allomas", {})
            robotok = m.get("robotok", [])
            self.kulcs_ertek([
                ("Szoftver", "RoboDK " + al.get("robodk_verzio", "?")),
                ("Állomás fájl", al.get("fajl") or al.get("nev", "")),
                ("Mérés időpontja", m.get("letrehozva", "")),
                ("Mérés megnevezése", m.get("megnevezes") or "–"),
                ("Robot(ok)", ", ".join("%s (%d tengely)" % (r["nev"], r.get("szabadsagfok", 0)) for r in robotok) or "–"),
                ("Szerszám(ok)", ", ".join(s["nev"] for s in m.get("szerszamok", [])) or "–"),
                ("Programok száma", str(len(m.get("programok", [])))),
                ("Célpontok száma", str(len(m.get("celpontok", [])))),
                ("Ütközésvizsgálat a mérés során", "igen" if m.get("beallitasok", {}).get("utkozesvizsgalat") else "nem"),
            ], "A szimulációs környezet adatai")
            for r in robotok:
                if r.get("also_hatar"):
                    self.tablazat(
                        ["Csukló", "Alsó határ [°]", "Felső határ [°]", "Teljes tartomány [°]", "Pillanatnyi érték [°]"],
                        [["J%d" % (k + 1), sz(r["also_hatar"][k]), sz(r["felso_hatar"][k]),
                          sz(r["felso_hatar"][k] - r["also_hatar"][k]),
                          sz(r["csuklok_fok"][k], 2) if k < len(r.get("csuklok_fok", [])) else "–"]
                         for k in range(len(r["also_hatar"]))],
                        "%s – csuklóhatárok" % r["nev"], szelessegek=[2.2, 3.3, 3.3, 3.6, 3.6], jobbra=[1, 2, 3, 4])

        # 3. állomás felépítése
        self.cimsor("%d.3 Az állomás (robotcella) felépítése" % fsz, 2)
        if m.get("kepek", {}).get("allomas"):
            self.abra(os.path.join(mm, m["kepek"]["allomas"]), "Az állomás a mérés időpontjában (RoboDK 3D nézet)")
        else:
            self.kitoltendo("illessz be egy képernyőképet a teljes állomásról.")
        self.kitoltendo("írd le röviden a cella felépítését: robot típusa, szerszám, referencia keretek és "
                        "objektumok elhelyezése, és miért így választottad.")
        poz_fejlec = ["Név", "Szülő", "X [mm]", "Y [mm]", "Z [mm]", "Rx [°]", "Ry [°]", "Rz [°]"]
        poz_szel = [3.4, 3.0, 1.6, 1.6, 1.6, 1.6, 1.6, 1.6]
        if m.get("referenciak"):
            self.tablazat(poz_fejlec, [[e["nev"], e["szulo"]] + [sz(v) for v in e["relativ"]] for e in m["referenciak"]],
                          "Referencia keretek (a szülőhöz képest)", poz_szel, jobbra=range(2, 8))
        if m.get("szerszamok"):
            self.tablazat(["Név", "Robot"] + poz_fejlec[2:], [[e["nev"], e["szulo"]] + [sz(v) for v in e["tcp"]]
                                                               for e in m["szerszamok"]],
                          "Szerszámok – TCP a robot karimájához (flange) képest", poz_szel, jobbra=range(2, 8))
        if m.get("objektumok"):
            self.tablazat(poz_fejlec, [[e["nev"], e["szulo"]] + [sz(v) for v in e["relativ"]] for e in m["objektumok"]],
                          "Objektumok (a szülőhöz képest)", poz_szel, jobbra=range(2, 8))

        # 4. lépések
        self.cimsor("%d.4 A megvalósítás lépései" % fsz, 2)
        if not f["lepesek"]:
            self.kitoltendo("írd le a megvalósítás lépéseit képernyőképekkel – vagy rögzítsd őket a RoboDK-ban a "
                            "lepes_rogzitese.py szkripttel, és futtasd újra a jegyzőkönyv-készítőt.")
        for n, l in enumerate(f["lepesek"], 1):  # folyamatos számozás, akkor is, ha töröltél lépést
            p = self.bekezdes(utana=3)
            p.paragraph_format.keep_with_next = True
            p.add_run("%d. lépés: " % n).bold = True
            p.add_run(l["leiras"])
            if l.get("kep"):
                self.abra(os.path.join(f["mappa"], l["kep"]), "%d. lépés: %s" % (n, l["leiras"]), 14.5)
            info = "Rögzítve: %s." % l.get("idopont", "")
            if l.get("uj_elemek"):
                info += " Új elemek az állomásban: %s." % ", ".join(l["uj_elemek"])
            self.megjegyzes(info)

        # 5. célpontok
        self.cimsor("%d.5 Célpontok (targetek)" % fsz, 2)
        celpontok = m.get("celpontok", [])
        if not celpontok:
            self.megjegyzes("Az állomásban nincs célpont, vagy nem készült mérés.")
        else:
            self.tablazat(
                ["Név", "Referencia", "Típus", "X [mm]", "Y [mm]", "Z [mm]", "Rx [°]", "Ry [°]", "Rz [°]"],
                [[c["nev"], c["referencia"], "csukló" if c.get("csuklo_target") else "Descartes"] + [sz(v) for v in c["poz"]]
                 for c in celpontok],
                "Célpontok helyzete és orientációja (a referencia kerethez képest)",
                [3.1, 2.4, 1.8, 1.5, 1.5, 1.5, 1.4, 1.4, 1.4], jobbra=range(3, 9), meret=self.args.betumeret - 2.5)
            n = max(len(c.get("csuklok_fok", [])) for c in celpontok)
            if n:
                self.tablazat(
                    ["Név"] + ["J%d [°]" % (k + 1) for k in range(n)],
                    [[c["nev"]] + [sz(v, 2) for v in c.get("csuklok_fok", [])] + ["–"] * (n - len(c.get("csuklok_fok", [])))
                     for c in celpontok],
                    "Célpontok csuklószögei (a célpontban felvett robotkonfiguráció)",
                    [3.4] + [12.6 / n] * n, jobbra=range(1, n + 1), meret=self.args.betumeret - 2.5)

        # 6. programok + mérések
        self.cimsor("%d.6 Programok és mérési eredmények" % fsz, 2)
        if len(f["osszes_meres"]) > 1:
            self.megjegyzes("A részletes eredmények a legutolsó mérésre vonatkoznak (%s, %s); az összes mérés "
                            "összehasonlítása a következő alfejezetben található." % (
                                m.get("megnevezes") or "megnevezés nélkül", m.get("letrehozva", "")))
        programok = m.get("programok", [])
        if not programok:
            self.megjegyzes("Nem található program az állomásban.")
        self._alairasok, self._programok, self._abrazolt = {}, programok, set()
        for pi, p in enumerate(programok, 1):
            self.program(fsz, pi, p, mm)

        alfejezet = 7
        if len([p for p in programok if p.get("frissites")]) > 1:
            self.programok_osszehasonlitasa(fsz, alfejezet, programok, mm)
            alfejezet += 1
        if len(f["osszes_meres"]) > 1:
            self.osszehasonlitas(fsz, alfejezet, f)
            alfejezet += 1

        # értékelés
        self.cimsor("%d.%d Értékelés" % (fsz, alfejezet), 2)
        if len(programok) == 1:
            self.bekezdes("A mérések alapján (automatikusan összeállított összefoglaló):", felkover=True)
            for mondat in automatikus_ertekeles(programok[0]):
                self.felsorolas(mondat)
        elif programok:
            self.bekezdes("A mérések alapján (automatikusan összeállított összefoglaló):", felkover=True)
            mondatok = [automatikus_ertekeles(p)[1:] for p in programok]  # az 1. mondat programonként egyedi
            kozos = [x for x in mondatok[0] if all(x in m_ for m_ in mondatok[1:])]
            if kozos:
                self.bekezdes("Minden programra érvényes:", utana=2).paragraph_format.keep_with_next = True
                for mondat in kozos:
                    self.felsorolas(mondat)
            self.bekezdes("Programonként:", utana=2).paragraph_format.keep_with_next = True
            for p, mm_ in zip(programok, mondatok):
                egyedi = [x for x in mm_ if x not in kozos]
                self.felsorolas(" ".join([automatikus_ertekeles(p)[0]] + egyedi))
        self.kitoltendo("saját értékelés: teljesíti-e a megoldás a feladat követelményeit? Megfelel-e a ciklusidő? "
                        "Mit lehetne javítani (sebesség, lekerekítés/blending, MoveJ a MoveL helyett, célpontok "
                        "áthelyezése)? Milyen problémák merültek fel (elérhetőség, szingularitás, ütközés), és hogyan "
                        "oldottad meg őket?")
        if m.get("figyelmeztetesek"):
            self.megjegyzes("A mérőszkript figyelmeztetései: " + " | ".join(m["figyelmeztetesek"]))

    def programok_osszehasonlitasa(self, fsz, alfejezet, programok, mm):
        self.cimsor("%d.%d Programok összehasonlítása" % (fsz, alfejezet), 2)
        progs = [p for p in programok if p.get("frissites")]
        st = {p["nev"]: (p.get("palya") or {}).get("statisztika") or {} for p in progs}
        van_zona = any(s.get("zona") for s in st.values())
        fejlec = ["Program", "Ciklusidő [s]", "Pályahossz [mm]", "Max TCP seb. [mm/s]", "Max eltérés, megálló [mm]",
                  "Max eltérés egyenestől [mm]"] + (["Zóna min. táv. [mm]"] if van_zona else [])
        sorok = []
        for p in progs:
            s = st[p["nev"]]
            sor = [p["nev"], sz(p["frissites"]["ciklusido_s"], 3), sz(p["frissites"]["palyahossz_mm"], 1),
                   sz(s.get("tcp_max_sebesseg_mm_s"), 1), sz(s.get("max_elteres_megallo_mm"), 4),
                   sz(s.get("max_palya_elteres_mm"), 3)]
            if van_zona:
                sor.append(sz((s.get("zona") or {}).get("min_tavolsag_mm"), 1))
            sorok.append(sor)
        szel = [4.2, 1.9, 2.0, 2.0, 2.0, 2.0] + ([1.9] if van_zona else [])
        szel = [w * 16.0 / sum(szel) for w in szel]
        self.tablazat(fejlec, sorok, "A programok fő mérési eredményei", szel, jobbra=range(1, len(fejlec)),
                      meret=self.args.betumeret - 2.5)
        abrak = os.path.join(mm, "abrak")
        os.makedirs(abrak, exist_ok=True)
        abra_program_idok([p["nev"] for p in progs], [p["frissites"]["ciklusido_s"] for p in progs],
                          os.path.join(abrak, "programok_ciklusido.png"))
        self.abra(os.path.join(abrak, "programok_ciklusido.png"), "A programok ciklusideje")
        parok = movej_movel_parok(progs)
        if parok:
            psorok = []
            for valtozat, j, l in parok:
                tj, tl = j["frissites"]["ciklusido_s"], l["frissites"]["ciklusido_s"]
                psorok.append([valtozat, sz(tj, 3), sz(tl, 3), sz(tl - tj, 3), szazalek(100.0 * (tl - tj) / tj if tj else None, 1),
                               sz(j["frissites"]["palyahossz_mm"], 1), sz(l["frissites"]["palyahossz_mm"], 1)])
            self.tablazat(["Változat", "MoveJ [s]", "MoveL [s]", "Különbség [s]", "Különbség [%]", "Út MoveJ [mm]",
                           "Út MoveL [mm]"], psorok, "A MoveJ és a MoveL változat végrehajtási ideje és úthossza",
                          [2.8, 2.0, 2.0, 2.2, 2.2, 2.4, 2.4], jobbra=range(1, 7))
            self.megjegyzes("Különbség = MoveL − MoveJ, a százalék a MoveJ idejéhez viszonyítva; pozitív érték "
                            "esetén a MoveJ-s változat a gyorsabb.")
            abra_movej_movel([v for v, _, _ in parok], [j["frissites"]["ciklusido_s"] for _, j, _ in parok],
                             [l["frissites"]["ciklusido_s"] for _, _, l in parok], os.path.join(abrak, "movej_movel.png"))
            self.abra(os.path.join(abrak, "movej_movel.png"), "A MoveJ és a MoveL változat ciklusideje")
        adatok = []
        for p in progs:
            try:
                adatok.append((p["nev"], palya_csv(os.path.join(mm, p["palya"]["csv"]))))
            except Exception:
                pass
        if len(adatok) > 1:
            abra_kis_tobbszoros(adatok, "ido_s", "TCP_sebesseg_mm_s", "idő [s]", "TCP seb. [mm/s]",
                                os.path.join(abrak, "kis_tcp_sebesseg.png"))
            self.abra(os.path.join(abrak, "kis_tcp_sebesseg.png"),
                      "A TCP pályasebessége programonként (közös függőleges skála)")
            abra_kis_tobbszoros(adatok, "TCP_X_mm", "TCP_Y_mm", "X [mm]", "Y [mm]",
                                os.path.join(abrak, "kis_tcp_palya.png"), felulnezet=True)
            self.abra(os.path.join(abrak, "kis_tcp_palya.png"), "A TCP pályája felülnézetben programonként (azonos skála)")
            zonas = [(nev, d) for nev, d in adatok if "zona_tavolsag_mm" in d]
            if zonas:
                abra_kis_tobbszoros(zonas, "ido_s", "zona_tavolsag_mm", "idő [s]", "táv. a zónától [mm]",
                                    os.path.join(abrak, "kis_zona.png"), zona=True)
                self.abra(os.path.join(abrak, "kis_zona.png"), "A TCP távolsága a tiltott zónától programonként")
        self.bekezdes("Megállapítások (a mérésekből automatikusan):", felkover=True)
        for mondat in osszehasonlito_mondatok(progs):
            self.felsorolas(mondat)

    def osszehasonlitas(self, fsz, alfejezet, f):
        self.cimsor("%d.%d Mérések összehasonlítása" % (fsz, alfejezet), 2)
        self.bekezdes("A feladathoz %d mérés készült. Az alábbi táblázat mindegyik mérés fő eredményeit tartalmazza "
                      "(a mérések megnevezését a mérőszkript indításakor lehet megadni)." % len(f["osszes_meres"]))
        sorok, programonkent = [], {}
        for i, mr in enumerate(f["osszes_meres"], 1):
            a = mr["adat"]
            cimke = a.get("megnevezes") or "%d. mérés" % i
            for p in a.get("programok", []):
                fr, ut = p.get("frissites") or {}, p.get("utkozes")
                st = (p.get("palya") or {}).get("statisztika") or {}
                sorok.append([cimke + "\n(" + a.get("letrehozva", "")[:16] + ")", p["nev"], sz(fr.get("ciklusido_s"), 2),
                              sz(fr.get("palyahossz_mm"), 1), sz(st.get("tcp_max_sebesseg_mm_s"), 1),
                              szazalek(100.0 * fr["ervenyesseg_arany"], 0) if fr else "–",
                              "–" if not ut else ("nincs" if ut.get("utkozesmentes") else "VAN")])
                if fr.get("ciklusido_s") is not None:
                    programonkent.setdefault(p["nev"], []).append((cimke, fr["ciklusido_s"]))
        self.tablazat(["Mérés", "Program", "Ciklusidő [s]", "Pályahossz [mm]", "Max TCP seb. [mm/s]",
                       "Érvényes", "Ütközés"], sorok, "Az összes mérés fő eredményei",
                      [3.8, 2.9, 1.8, 2.0, 2.0, 1.9, 1.6], jobbra=[2, 3, 4, 5])
        if programonkent:
            fajl = os.path.join(f["meres_mappa"], "abrak", "osszehasonlitas.png")
            os.makedirs(os.path.dirname(fajl), exist_ok=True)
            abra_osszehasonlitas(list(programonkent.items()), fajl)
            self.abra(fajl, "A ciklusidő mérésenként és programonként")
        self.kitoltendo("értelmezd az eltéréseket: melyik paraméter (sebesség, lekerekítés, mozgástípus, célpontok "
                        "helye) mennyivel változtatta a ciklusidőt és a pályát?")

    def program(self, fsz, pi, p, mm):
        self.cimsor("%d.6.%d %s" % (fsz, pi, p["nev"]), 3)
        if p.get("robot"):
            self.bekezdes("Robot: %s. Utasítások száma: %d." % (p["robot"], len(p.get("utasitasok", []))))
        alairas = program_alairas(p)
        azonos = self._alairasok.get(alairas)
        if azonos and p.get("utasitasok"):
            masik = next(q for q in self._programok if q["nev"] == azonos)
            elteres = [u["nev"] for u, v in zip(p["utasitasok"], masik["utasitasok"]) if u["nev"] != v["nev"]]
            self.bekezdes("Az utasítások szerkezete megegyezik a(z) %s programéval%s." % (
                azonos, (", eltérés: " + "; ".join(elteres)) if elteres else ""))
        elif p.get("utasitasok"):
            self._alairasok[alairas] = p["nev"]
            self.tablazat(
                ["#", "Utasítás (RoboDK)", "Típus", "Mozgás", "Cél"],
                [[u["sorszam"], u["nev"], u["tipus"], u.get("mozgas", ""),
                  {True: "csukló", False: "Descartes"}.get(u.get("csuklo_cel"), "")] for u in p["utasitasok"]],
                "A(z) %s program utasításai" % p["nev"], [1.0, 6.4, 3.4, 3.2, 2.0], jobbra=[0])

        fr, ut = p.get("frissites"), p.get("utkozes")
        st = (p.get("palya") or {}).get("statisztika")
        if not fr:
            self.megjegyzes("A program nem volt ellenőrizhető (lásd a figyelmeztetéseket).")
            return
        sorok = [
            ("Becsült ciklusidő (RoboDK)", sz(fr["ciklusido_s"], 3, "s")),
            ("TCP pályahossz (RoboDK)", sz(fr["palyahossz_mm"], 1, "mm")),
            ("Végrehajtható rész (érvényesség)", szazalek(100.0 * fr["ervenyesseg_arany"])),
            ("Érvényes utasítások", "%d / %d" % (fr["ervenyes_utasitasok"], len(p.get("utasitasok", [])))),
            ("RoboDK üzenet", fr.get("uzenet") or "–"),
        ]
        if ut:
            sorok.append(("Ütközésvizsgálat", "nincs ütközés" if ut.get("utkozesmentes")
                          else "ütközés vagy egyéb hiba (érvényesség: %s) – %s" % (szazalek(100.0 * ut["ervenyesseg_arany"]), ut.get("uzenet", ""))))
        if st:
            sorok += [
                ("Ciklusidő a mintavételezett pályából", sz(st["idotartam_s"], 3, "s")),
                ("TCP pályahossz a mintavételezett pályából", sz(st["tcp_palyahossz_mm"], 1, "mm")),
                ("Átlagos TCP sebesség", sz(st["tcp_atlag_sebesseg_mm_s"], 1, "mm/s")),
                ("Legnagyobb TCP sebesség", sz(st["tcp_max_sebesseg_mm_s"], 1, "mm/s")),
            ]
            if st.get("felveteltol_lerakasig_s") is not None:
                sorok.append(("Felvételtől lerakásig (a felvételi pont elhagyásától a lerakási pont eléréséig)",
                              sz(st["felveteltol_lerakasig_s"], 3, "s")))
            if st.get("max_elteres_megallo_mm") is not None:
                sorok.append(("Legnagyobb eltérés a megálló célpontokban", sz(st["max_elteres_megallo_mm"], 4, "mm")))
            if st.get("max_elteres_lekerekitett_mm") is not None:
                sorok.append(("Legnagyobb eltérés a lekerekített (nem megálló) célpontokban",
                              sz(st["max_elteres_lekerekitett_mm"], 3, "mm")))
            if st.get("max_palya_elteres_mm") is not None:
                sorok.append(("Legnagyobb eltérés az egyenes pályától (MoveL szakaszok)", sz(st["max_palya_elteres_mm"], 3, "mm")))
            if st.get("ismetlesi_elteres_mm") is not None:
                sorok.append(("Ugyanazon célpont ismételt elérésének eltérése", sz(st["ismetlesi_elteres_mm"], 4, "mm")))
            if st.get("kezdo_veg_elteres_mm") is not None:
                sorok.append(("A program kezdő- és végpontjának távolsága", sz(st["kezdo_veg_elteres_mm"], 3, "mm")))
            z = st.get("zona")
            if z:
                sorok.append(("Legkisebb távolság a tiltott zónától (TCP)",
                              ("BELÉPETT a zónába (%s)" % z["zona"]) if z["belepett"] else "%s (t = %s, %s)" % (
                                  sz(z["min_tavolsag_mm"], 1, "mm"), sz(z["ido_s"], 2, "s"), z["zona"])))
            sorok += [
                ("Mintavétel", "%d minta, %s időlépés" % (st["mintak_szama"], sz(p["palya"]["idolepes_s"], 4, "s"))),
                ("Hibás minták (szingularitás, elérhetőség, ütközés)", str(st["hibas_mintak"])),
            ]
        self.kulcs_ertek(sorok, "A(z) %s program fő mérési eredményei" % p["nev"])

        palya = p.get("palya")
        if not palya:
            return
        cp = palya.get("celpontok")
        if cp:
            self.tablazat(
                ["#", "Célpont", "Mozgás", "Érkezés [s]", "Szakaszidő [s]", "Úthossz [mm]", "Eltérés [mm]", "Egyenestől [mm]"],
                [[c["sorszam"], c["celpont"] + ("" if c["megallt"] else " *"), c["mozgas"].split(" ")[0],
                  sz(c["erkezes_s"], 3), sz(c["szakaszido_s"], 3), sz(c["szakasz_hossz_mm"], 1), sz(c["elteres_mm"], 4),
                  sz(c["palya_elteres_mm"], 3) if "palya_elteres_mm" in c else "–"] for c in cp],
                "Célpontok elérése és pályaszakaszok – %s" % p["nev"], [0.8, 2.8, 1.6, 1.9, 2.1, 2.0, 2.2, 2.6],
                jobbra=[0, 3, 4, 5, 6, 7])
            self.megjegyzes("Szakaszidő és úthossz: az előző célpont elhagyásától ennek eléréséig. Eltérés: a célpont "
                            "és a pálya legkisebb távolsága (megálló célpontnál a mintavétel felbontásán belüli érték). "
                            "Egyenestől: MoveL-szakaszon a pálya legnagyobb távolsága az ideális egyenestől. "
                            "* = lekerekített, nem megálló célpont.")
            n_cs = len(cp[0]["csuklo_valtozas_fok"])
            osszes = [sum(abs(c["csuklo_valtozas_fok"][k]) for c in cp) for k in range(n_cs)]
        if cp and (alairas not in self._abrazolt or self.args.reszletes):  # azonos célpontoknál a tengelyelfordulás is azonos
            self.tablazat(
                ["#", "Célpont"] + ["ΔJ%d [°]" % (k + 1) for k in range(n_cs)],
                [[c["sorszam"], c["celpont"]] + [sz(v, 1) for v in c["csuklo_valtozas_fok"]] for c in cp]
                + [["", "Σ |Δ|"] + [sz(v, 1) for v in osszes]],
                "A tengelyek elfordulása szakaszonként – %s" % p["nev"], [0.8, 2.8] + [12.4 / n_cs] * n_cs,
                jobbra=[0] + list(range(2, n_cs + 2)), meret=self.args.betumeret - 2.5)
        kepviselo = alairas not in self._abrazolt
        self._abrazolt.add(alairas)
        if not (kepviselo or self.args.reszletes):
            self.megjegyzes("A pályakép és a grafikonok a(z) %s programéhoz hasonlók; a programokat a %d.7 alfejezet "
                            "közös skálájú grafikonjai hasonlítják össze." % (self._alairasok.get(alairas, "előző"), fsz))
        if palya.get("kep") and (kepviselo or self.args.reszletes):
            self.abra(os.path.join(mm, palya["kep"]), "A(z) %s program pályája (átlátszó „szellemrobotok”)" % p["nev"])
        try:
            d = palya_csv(os.path.join(mm, palya["csv"]))
        except Exception as e:
            self.megjegyzes("A pályaadatok nem olvashatók (%s)." % e)
            d = None
        abrak = os.path.join(mm, "abrak")
        os.makedirs(abrak, exist_ok=True)
        alap = os.path.join(abrak, os.path.splitext(palya["csv"])[0])
        n = len(st["csuklok"]) if st else 0
        if d and n and (kepviselo or self.args.reszletes):
            abra_csuklo_idosor(d, n, "J%d_fok", "°", alap + "_csuklok.png")
            self.abra(alap + "_csuklok.png", "Csuklószögek az idő függvényében – %s" % p["nev"])
            if "v_J1_fok_s" in d and self.args.reszletes:
                abra_csuklo_idosor(d, n, "v_J%d_fok_s", "°/s", alap + "_csuklosebesseg.png")
                self.abra(alap + "_csuklosebesseg.png", "Csuklók szögsebessége az idő függvényében – %s" % p["nev"])
            abra_tcp_sebesseg(d, alap + "_tcp_sebesseg.png")
            self.abra(alap + "_tcp_sebesseg.png", "A TCP pályasebessége az idő függvényében – %s" % p["nev"])
            abra_tcp_palya(d, alap + "_tcp_palya.png")
            self.abra(alap + "_tcp_palya.png", "A TCP pályája felül- és oldalnézetben – %s" % p["nev"])
            if "zona_tavolsag_mm" in d:
                abra_zona_tavolsag(d, alap + "_zona.png")
                self.abra(alap + "_zona.png", "A TCP távolsága a tiltott zónától – %s" % p["nev"])
        if st:
            self.tablazat(
                ["Csukló", "Min [°]", "Max [°]", "Bejárt [°]", "Tartalék alsó / felső [°]", "Kihaszn.", "Max ω [°/s]", "Max ε [°/s²]"],
                [[c["nev"], sz(c["min"]), sz(c["max"]), sz(c["mozgastartomany"]),
                  ("%s / %s" % (sz(c["tartalek_also"]), sz(c["tartalek_felso"]))) if "tartalek_also" in c else "–",
                  szazalek(c.get("kihasznaltsag_szazalek"), 0), sz(c.get("max_sebesseg_fok_s")), sz(c.get("max_gyorsulas_fok_s2"), 0)]
                 for c in st["csuklok"]],
                "Csuklóstatisztika – %s" % p["nev"], [1.5, 1.7, 1.7, 1.8, 3.3, 1.6, 2.1, 2.3], jobbra=range(1, 8))
            if self.args.reszletes and abra_kihasznaltsag(st["csuklok"], alap + "_kihasznaltsag.png"):
                self.abra(alap + "_kihasznaltsag.png", "A csuklótartományok kihasználtsága – %s" % p["nev"])
            szak = st.get("szakaszok", []) if not cp else []
            if szak:
                max_sor = 60
                self.tablazat(
                    ["#", "MOVE_ID", "Kezdet [s]", "Vége [s]", "Időtartam [s]", "Úthossz [mm]"],
                    [[s["sorszam"], s["move_id"], sz(s["kezdet_s"], 3), sz(s["veg_s"], 3), sz(s["idotartam_s"], 3),
                      sz(s["hossz_mm"], 1)] for s in szak[:max_sor]],
                    "Mozgásszakaszok végrehajtási sorrendben – %s" % p["nev"], [1.2, 2.2, 3.0, 3.0, 3.2, 3.4], jobbra=range(0, 6))
                if len(szak) > max_sor:
                    self.megjegyzes("A táblázat az első %d szakaszt mutatja; a teljes lista a meresek.json fájlban van." % max_sor)
                self.megjegyzes("A szakaszokat a RoboDK pályaszámítás mozgásazonosítója (MOVE_ID) választja el; "
                                "sorrendjük megegyezik a program mozgásutasításainak sorrendjével.")
            if st.get("hibak"):
                self.tablazat(["Hiba", "Első előfordulás [s]", "Hibás minták"],
                              [[h["leiras"], sz(h["elso_idopont_s"], 3), h["mintak"]] for h in st["hibak"]],
                              "A pályán észlelt hibák – %s" % p["nev"], [10.0, 3.2, 2.8], jobbra=[1, 2])

    def osszefoglalas(self, feladatok):
        self.uj_oldal()
        self.cimsor("Összefoglalás", 1)
        sorok = []
        for fsz, f in enumerate(feladatok, 1):
            for p in (f["meres"] or {}).get("programok", []):
                fr = p.get("frissites") or {}
                ut = p.get("utkozes")
                st = (p.get("palya") or {}).get("statisztika") or {}
                sorok.append(["%d." % fsz, p["nev"], sz(fr.get("ciklusido_s"), 2), sz(fr.get("palyahossz_mm"), 1),
                              sz(st.get("tcp_max_sebesseg_mm_s"), 1),
                              szazalek(100.0 * fr["ervenyesseg_arany"], 0) if fr else "–",
                              "–" if not ut else ("nincs" if ut.get("utkozesmentes") else "VAN")])
        if sorok:
            self.tablazat(["Feladat", "Program", "Ciklusidő [s]", "Pályahossz [mm]", "Max TCP seb. [mm/s]", "Érvényes",
                           "Ütközés"], sorok, "Az összes program fő mérési eredményei",
                          [1.4, 4.4, 2.0, 2.2, 2.2, 1.9, 1.9], jobbra=[2, 3, 4, 5])
        self.kitoltendo("néhány mondatos összefoglalás: mit valósítottál meg, mik a legfontosabb eredmények, mit tanultál.")

        self.cimsor("Mellékletek", 1)
        self.bekezdes("A jegyzőkönyv táblázatai és grafikonjai az alábbi, a RoboDK-ból automatikusan kimentett "
                      "adatfájlokon alapulnak (pontosvesszővel tagolt CSV, Excelben megnyitható):")
        for fsz, f in enumerate(feladatok, 1):
            reszek = []
            if f["lepesek"]:
                reszek.append("lépésnapló: lepesek.json és kepek/")
            if f["meres_mappa"]:
                reszek.append("mérések: %s/ (%s)" % (os.path.basename(f["meres_mappa"]), ", ".join(
                    x for x in sorted(os.listdir(f["meres_mappa"])) if x.endswith((".csv", ".json")))))
            self.felsorolas("%d. feladat – %s mappa: %s" % (
                fsz, os.path.basename(f["mappa"].rstrip("/\\")), "; ".join(reszek) or "nincs adat"))


def automatikus_ertekeles(p):
    """Tényszerű, a mérésekből levezetett mondatok egy programról."""
    fr, ut = p.get("frissites"), p.get("utkozes")
    st = (p.get("palya") or {}).get("statisztika")
    if not fr:
        return ["A(z) %s program nem volt ellenőrizhető." % p["nev"]]
    m = ["A(z) %s program %d utasításból áll; a RoboDK által becsült ciklusidő %s, a TCP által megtett út %s."
         % (p["nev"], len(p.get("utasitasok", [])), sz(fr["ciklusido_s"], 2, "s"), sz(fr["palyahossz_mm"], 1, "mm"))]
    if fr["ervenyesseg_arany"] >= 0.999:
        m.append("A pálya teljes egészében végrehajtható (érvényesség: 100%).")
    else:
        m.append("A pálya csak %s-ban hajtható végre – RoboDK üzenet: „%s”."
                 % (szazalek(100.0 * fr["ervenyesseg_arany"]), fr.get("uzenet", "")))
    if ut:
        m.append("Az ütközésvizsgálat nem jelzett ütközést." if ut.get("utkozesmentes")
                 else "Az ütközésvizsgálattal futtatott ellenőrzés ütközést vagy egyéb hibát jelzett (%s)." % ut.get("uzenet", ""))
    if st:
        m.append("Az átlagos TCP-sebesség %s, a legnagyobb %s."
                 % (sz(st["tcp_atlag_sebesseg_mm_s"], 1, "mm/s"), sz(st["tcp_max_sebesseg_mm_s"], 1, "mm/s")))
        cs = [c for c in st["csuklok"] if c.get("kihasznaltsag_szazalek") is not None]
        if cs:
            legtobb = max(cs, key=lambda c: c["kihasznaltsag_szazalek"])
            legkozelebb = min(cs, key=lambda c: min(c["tartalek_also"], c["tartalek_felso"]))
            m.append("A csuklók legfeljebb a mozgástartományuk %s-át használták ki (%s); a csuklóhatárhoz legközelebb "
                     "a(z) %s csukló került (%s tartalék)." % (
                         szazalek(legtobb["kihasznaltsag_szazalek"], 0), legtobb["nev"], legkozelebb["nev"],
                         sz(min(legkozelebb["tartalek_also"], legkozelebb["tartalek_felso"]), 1, "°")))
        if st.get("felveteltol_lerakasig_s") is not None:
            m.append("A felvételi pont elhagyásától a lerakási pont eléréséig %s telt el."
                     % sz(st["felveteltol_lerakasig_s"], 2, "s"))
        cp = (p.get("palya") or {}).get("celpontok") or []
        megallo = [c for c in cp if c["megallt"]]
        if megallo:
            felb = max(c["felbontas_mm"] for c in megallo)
            if st["max_elteres_megallo_mm"] <= max(felb, 0.001):
                m.append("A %d megálló célpontot a robot a mintavétel felbontásán belül (≤ %s) érte el, vagyis "
                         "gyakorlatilag eltérés nélkül." % (len(megallo), sz(max(felb, 0.001), 3, "mm")))
            else:
                m.append("A megálló célpontokban a legnagyobb eltérés %s volt." % sz(st["max_elteres_megallo_mm"], 3, "mm"))
        if st.get("max_elteres_lekerekitett_mm") is not None:
            m.append("A lekerekített célpontoknál a pálya legfeljebb %s-re közelítette meg a célpontot: ezt a "
                     "lekerekítés (blending) okozza." % sz(st["max_elteres_lekerekitett_mm"], 2, "mm"))
        if st.get("max_palya_elteres_mm") is not None:
            if st["max_palya_elteres_mm"] < 0.001:
                m.append("A MoveL-szakaszokon a pálya mérhetően nem tért el az ideális egyenestől (< 0,001 mm).")
            else:
                m.append("A MoveL-szakaszokon a pálya legfeljebb %s-re tért el az ideális egyenestől."
                         % sz(st["max_palya_elteres_mm"], 3, "mm"))
        if st.get("ismetlesi_elteres_mm") is not None:
            m.append("Ugyanazt a célpontot többször elérve az elért helyzetek legfeljebb %s-re estek egymástól "
                     "(ismételhetőség a szimulációban)." % sz(st["ismetlesi_elteres_mm"], 4, "mm"))
        z = st.get("zona")
        if z:
            m.append("A TCP BELÉPETT a tiltott zónába (%s)!" % z["zona"] if z["belepett"] else
                     "A TCP legkisebb távolsága a tiltott zónától %s volt (t = %s): a robot nem lépett be a zónába."
                     % (sz(z["min_tavolsag_mm"], 1, "mm"), sz(z["ido_s"], 2, "s")))
        if st["hibas_mintak"]:
            m.append("A mintavételezett pályán %d mintában jelentkezett hiba: %s."
                     % (st["hibas_mintak"], "; ".join(h["leiras"] for h in st["hibak"])))
        else:
            m.append("A mintavételezett pályán nem jelentkezett szingularitás, elérhetetlenség vagy ütközés.")
    return m


def osszehasonlito_mondatok(progs):
    """Több program összevetéséből levont, tényszerű megállapítások."""
    m = []
    idok = [(p["frissites"]["ciklusido_s"], p["nev"]) for p in progs]
    (tmin, gyors), (tmax, lassu) = min(idok), max(idok)
    m.append("A leggyorsabb program a(z) %s (%s), a leglassabb a(z) %s (%s); a különbség %s."
             % (gyors, sz(tmin, 2, "s"), lassu, sz(tmax, 2, "s"), sz(tmax - tmin, 2, "s")))
    for valtozat, j, l in movej_movel_parok(progs):
        tj, tl = j["frissites"]["ciklusido_s"], l["frissites"]["ciklusido_s"]
        lj, ll = j["frissites"]["palyahossz_mm"], l["frissites"]["palyahossz_mm"]
        if tj <= tl:
            szoveg = "a MoveJ-es változat %s-mal gyorsabb (a MoveL-es %s-kal tovább tart)" % (
                sz(tl - tj, 2, "s"), szazalek(100.0 * (tl - tj) / tj if tj else None, 1))
        else:
            szoveg = "a MoveL-es változat %s-mal gyorsabb (a MoveJ-es %s-kal tovább tart)" % (
                sz(tj - tl, 2, "s"), szazalek(100.0 * (tj - tl) / tl if tl else None, 1))
        m.append("%s: %s; a TCP útja MoveJ-vel %s, MoveL-lel %s (a MoveL egyenes vonalon visz)." % (
            valtozat, szoveg, sz(lj, 0, "mm"), sz(ll, 0, "mm")))
    st = [((p.get("palya") or {}).get("statisztika") or {}, p) for p in progs]
    megallo = [s["max_elteres_megallo_mm"] for s, _ in st if s.get("max_elteres_megallo_mm") is not None]
    if megallo:
        m.append("A megálló célpontokban mért legnagyobb eltérés minden programban legfeljebb %s volt: a sebesség "
                 "a célpontok elérésének pontosságát a szimulációban nem befolyásolta, csak a végrehajtási időt."
                 % sz(max(megallo), 3, "mm"))
    kerek = [(s["max_elteres_lekerekitett_mm"], p["nev"]) for s, p in st if s.get("max_elteres_lekerekitett_mm") is not None]
    if kerek:
        m.append("Lekerekítéssel (%s) a célpontok közelében legfeljebb %s eltérés jelent meg, cserébe rövidebb a "
                 "ciklusidő. A pontosság és a sebesség között tehát a lekerekítés jelent kompromisszumot, nem maga "
                 "a sebesség." % (", ".join(n for _, n in kerek), sz(max(v for v, _ in kerek), 2, "mm")))
    zonak = [s["zona"]["min_tavolsag_mm"] for s, _ in st if s.get("zona")]
    if zonak:
        belepett = [p["nev"] for s, p in st if s.get("zona") and s["zona"]["belepett"]]
        m.append(("A TCP a következő programokban belépett a tiltott zónába: %s." % ", ".join(belepett)) if belepett else
                 "A TCP és a tiltott zóna legkisebb távolsága a programokban %s és %s között volt; a robot egyik "
                 "változatban sem lépett be a zónába." % (sz(min(zonak), 1, "mm"), sz(max(zonak), 1, "mm")))
    return m


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------
def pdf_keszites(docx_fajl):
    pdf_fajl = os.path.splitext(docx_fajl)[0] + ".pdf"
    try:  # Microsoft Word (Windows / macOS)
        from docx2pdf import convert
        convert(docx_fajl, pdf_fajl)
        if os.path.isfile(pdf_fajl):
            return pdf_fajl
    except Exception:
        pass
    jeloltek = [shutil.which("soffice"), shutil.which("libreoffice"),
                r"C:\Program Files\LibreOffice\program\soffice.exe", "/Applications/LibreOffice.app/Contents/MacOS/soffice"]
    for s in jeloltek:
        if s and os.path.isfile(s):
            # külön profil: akkor is működik, ha a LibreOffice éppen nyitva van
            with tempfile.TemporaryDirectory() as profil:
                subprocess.run([s, "-env:UserInstallation=" + pathlib.Path(profil).as_uri(), "--headless",
                                "--convert-to", "pdf", "--outdir", os.path.dirname(os.path.abspath(docx_fajl)), docx_fajl],
                               check=False, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=300)
            if os.path.isfile(pdf_fajl):
                return pdf_fajl
    return None


def main():
    ap = argparse.ArgumentParser(description="Word/PDF jegyzőkönyv készítése RoboDK-ban rögzített lépésekből és mérésekből.")
    ap.add_argument("mappak", nargs="+", help="a feladatok '<állomás>_dokumentacio' mappái, sorrendben")
    ap.add_argument("-o", "--kimenet", default="jegyzokonyv.docx", help="a Word fájl neve (alap: jegyzokonyv.docx)")
    ap.add_argument("--cim", default="RoboDK szimulációs feladatok", help="a jegyzőkönyv címe")
    ap.add_argument("--targy", default="[Tantárgy neve]", help="tantárgy")
    ap.add_argument("--intezmeny", default="[Egyetem, kar, tanszék]", help="intézmény")
    ap.add_argument("--nev", default="[Név]", help="hallgató neve")
    ap.add_argument("--neptun", default="", help="Neptun-kód")
    ap.add_argument("--oktato", default="", help="oktató neve")
    ap.add_argument("--betutipus", default="Calibri", help="betűtípus (pl. 'Times New Roman')")
    ap.add_argument("--betumeret", type=float, default=11, help="betűméret pontban")
    ap.add_argument("--pdf", action="store_true", help="PDF is készüljön (Word vagy LibreOffice kell hozzá)")
    ap.add_argument("--reszletes", action="store_true",
                    help="programonként a csuklósebesség- és csuklótartomány-grafikon is bekerül")
    ap.add_argument("--minta", action="store_true", help=argparse.SUPPRESS)
    args = ap.parse_args()

    feladatok = []
    for mappa in args.mappak:
        if not os.path.isdir(mappa):
            sys.exit("Nem létező mappa: %s" % mappa)
        feladatok.append(feladat_adatai(os.path.abspath(mappa)))

    jk = Jegyzokonyv(args)
    verzio = next((f["meres"]["allomas"].get("robodk_verzio") for f in feladatok if f["meres"]), "")
    jk.cimlap(verzio)
    jk.bevezetes(feladatok)
    for fsz, f in enumerate(feladatok, 1):
        print("%d. feladat: %s" % (fsz, f["cim"]))
        jk.feladat(fsz, f)
    jk.osszefoglalas(feladatok)
    jk.tartalomjegyzek()
    jk.fejlec_lablec()
    jk.doc.save(args.kimenet)
    print("Elkészült: " + os.path.abspath(args.kimenet))
    if args.pdf:
        pdf = pdf_keszites(os.path.abspath(args.kimenet))
        print("PDF: " + pdf if pdf else "PDF-et nem sikerült készíteni: nyisd meg Wordben -> Fájl -> Mentés másként -> PDF.")


if __name__ == "__main__":
    main()
